# document_ingestion_pipeline.py
# This script demonstrates setting up a CocoIndex data ingestion pipeline for internal documents.

# --- Installation Instructions for Text Extraction Libraries ---
# Before running a real version of this script, you'll need to install libraries for PDF and DOCX processing.
# You can install them using pip:
#
# pip install pypdf python-docx
#
# CocoIndex itself would also need to be installed:
# pip install cocoindex
#
# For this demonstration, CocoIndex and its modules are placeholders.
# ---------------------------------------------------------------

import os
from typing import Dict, Any, Iterator, List

# --- CocoIndex Placeholder Imports ---
# These are conceptual representations of what you might import from CocoIndex.
# In a real scenario, you would import these from the actual 'cocoindex' library.

class cocoindex:
    class sources:
        @staticmethod
        def LocalFile(path: str, file_filter: List[str] = None, recursive: bool = True):
            """
            Placeholder for LocalFile source.
            Reads files from a local directory.
            """
            print(f"[CocoIndex.sources.LocalFile] Initialized to read from: {path}, filter: {file_filter}, recursive: {recursive}")
            # Simulate finding files
            found_files = []
            for root, _, files in os.walk(path):
                for file in files:
                    if file_filter is None or any(file.endswith(ext) for ext in file_filter):
                        found_files.append(os.path.join(root, file))
            
            # Simulate yielding file content or file path dictionaries
            for filepath in found_files:
                try:
                    with open(filepath, 'r', encoding='utf-8') as f:
                        content = f.read()
                    yield {"path": filepath, "content": content, "metadata": {"source_filename": os.path.basename(filepath)}}
                except Exception as e:
                    # For binary files like PDF/DOCX, content might be binary or read differently
                    # This placeholder just notes it.
                    yield {"path": filepath, "content": f"Binary or special file: {filepath}", "metadata": {"source_filename": os.path.basename(filepath), "error": str(e)}}


    class functions:
        @staticmethod
        def SplitRecursively(text_field: str = "text_content", chunk_size: int = 1000, chunk_overlap: int = 200):
            """
            Placeholder for SplitRecursively function.
            Splits text into chunks.
            """
            print(f"[CocoIndex.functions.SplitRecursively] Initialized with text_field='{text_field}', chunk_size={chunk_size}, chunk_overlap={chunk_overlap}")
            
            def splitter(data_item: Dict[str, Any]) -> Iterator[Dict[str, Any]]:
                text_to_split = data_item.get(text_field, "")
                if not text_to_split:
                    yield data_item # Pass through if no text
                    return

                start = 0
                while start < len(text_to_split):
                    end = start + chunk_size
                    chunk = text_to_split[start:end]
                    new_item = data_item.copy()
                    new_item["chunk"] = chunk
                    new_item["metadata"] = new_item.get("metadata", {}).copy()
                    new_item["metadata"]["chunk_offset"] = start
                    yield new_item
                    if end >= len(text_to_split):
                        break
                    start += chunk_size - chunk_overlap
            return splitter

    @staticmethod
    def flow_def(name: str):
        """
        Placeholder for the flow_def decorator.
        """
        print(f"[CocoIndex.flow_def] Defining flow: {name}")
        def decorator(func):
            def wrapper(*args, **kwargs):
                print(f"[CocoIndex.flow_def] Running flow: {name}")
                return func(*args, **kwargs)
            return wrapper
        return decorator

    # Placeholder for a simple Flow execution context or runner
    class Flow:
        def __init__(self, source, *processors):
            self.source = source
            self.processors = processors
            print("[CocoIndex.Flow] Initialized with source and processors.")

        def run(self):
            print("[CocoIndex.Flow] Starting run...")
            processed_items = 0
            for item in self.source:
                current_item_or_items = item
                for processor_func in self.processors:
                    # Handle functions that might yield multiple items (like a splitter)
                    if isinstance(current_item_or_items, dict): # single item
                        current_item_or_items = list(processor_func(current_item_or_items))
                    else: # list of items
                        new_items_list = []
                        for sub_item in current_item_or_items:
                            new_items_list.extend(list(processor_func(sub_item)))
                        current_item_or_items = new_items_list
                
                # Simulate output/sink for demonstration
                if isinstance(current_item_or_items, list):
                    for final_item in current_item_or_items:
                        # print(f"  [Output Item]: {str(final_item)[:200]}...") # Print snippet of final item
                        processed_items +=1
                else: # Should ideally always be a list after splitter
                    # print(f"  [Output Item]: {str(current_item_or_items)[:200]}...")
                    processed_items +=1
            print(f"[CocoIndex.Flow] Run finished. Processed {processed_items} items (chunks).")
            # In a real scenario, this would return results or handle a sink.


# --- Text Extraction Functions ---

def extract_text_from_pdf_placeholder(file_path: str) -> str:
    """
    Placeholder for PDF text extraction.
    In a real implementation, you would use a library like 'pypdf'.
    Example:
        import pypdf
        reader = pypdf.PdfReader(file_path)
        text = ""
        for page in reader.pages:
            text += page.extract_text()
        return text
    """
    print(f"[Extractor] Reading PDF (placeholder): {file_path}")
    # Simulate reading some text
    return f"Extracted text from PDF: {os.path.basename(file_path)}. This is placeholder content."

def extract_text_from_docx_placeholder(file_path: str) -> str:
    """
    Placeholder for DOCX text extraction.
    In a real implementation, you would use a library like 'python-docx'.
    Example:
        from docx import Document
        doc = Document(file_path)
        text = ""
        for para in doc.paragraphs:
            text += para.text + "\\n"
        return text
    """
    print(f"[Extractor] Reading DOCX (placeholder): {file_path}")
    # Simulate reading some text
    return f"Extracted text from DOCX: {os.path.basename(file_path)}. This is placeholder content."

def extract_text_from_file(item: Dict[str, Any]) -> Dict[str, Any]:
    """
    Extracts text content from a file based on its extension.
    'item' is expected to be a dictionary with at least a 'path' key.
    The extracted text is added to 'item' under the key 'text_content'.
    """
    file_path = item.get("path")
    content = item.get("content") # From LocalFile source if it could read it

    if not file_path:
        item["text_content"] = ""
        item["metadata"]["extraction_error"] = "Missing file path"
        return item

    _, extension = os.path.splitext(file_path.lower())
    extracted_text = ""

    if extension == ".txt":
        print(f"[Extractor] Reading TXT: {file_path}")
        extracted_text = content # Use content directly if already read by source
    elif extension == ".md":
        print(f"[Extractor] Reading MD: {file_path}")
        extracted_text = content # Use content directly if already read by source
    elif extension == ".pdf":
        extracted_text = extract_text_from_pdf_placeholder(file_path)
    elif extension == ".docx":
        extracted_text = extract_text_from_docx_placeholder(file_path)
    else:
        print(f"[Extractor] Unsupported file type or using direct content: {file_path}")
        # Fallback or specific handling for other types, or rely on initial content
        extracted_text = content if content else f"Unsupported file type: {extension}"

    item["text_content"] = extracted_text
    if "content" in item: # remove raw content if text_content is populated
        del item["content"]
    return item

# --- CocoIndex Flow Definition ---

@cocoindex.flow_def(name="DocumentIngestionFlow")
def document_ingestion_flow():
    """
    Defines the data ingestion pipeline:
    1. Reads files from the './sample_documents/' directory.
    2. Extracts text content from these files.
    3. Splits the extracted text into manageable chunks.
    """
    # 1. Configure Data Source: Read from local files
    # It will look for .txt, .md, .pdf, .docx files in the specified path.
    # In a real CocoIndex setup, LocalFile would yield dictionaries with file info.
    file_source = cocoindex.sources.LocalFile(
        path="./sample_documents/",
        file_filter=[".txt", ".md", ".pdf", ".docx"],
        recursive=True
    )

    # 2. Define Processing Steps
    # - Text extraction is handled by the 'extract_text_from_file' function.
    #   This function will be mapped to each item from the source.
    # - Text splitting is handled by 'SplitRecursively'.
    #   It takes the 'text_content' field from the item and splits it.
    text_splitter = cocoindex.functions.SplitRecursively(
        text_field="text_content",  # Field containing the text to split
        chunk_size=500,            # Desired chunk size in characters
        chunk_overlap=50           # Overlap between chunks
    )

    # In a real CocoIndex flow, you would chain these using a Flow object or similar mechanism.
    # For this demonstration, we'll simulate the flow execution.
    # The flow would look something like:
    #
    #   flow = Flow(
    #       source=file_source,
    #       processors=[
    #           extract_text_from_file, # This would be a cocoindex.map or similar
    #           text_splitter
    #       ],
    #       sink=your_vector_store_sink  # Not covered in this script
    #   )
    #   flow.run()
    #
    # For this script, we'll use the placeholder Flow class to simulate.
    
    print("\n--- Simulating Document Ingestion Flow ---")
    # The actual CocoIndex framework would handle the iteration and passing of data.
    # Here, we manually simulate the chain for clarity using the placeholder Flow.
    
    # Note: In a real CocoIndex, `extract_text_from_file` would typically be wrapped
    # in something like `cocoindex.functions.Map(extract_text_from_file)` if it's not
    # implicitly handled by the flow runner. For simplicity, our placeholder Flow
    # will just call it.

    simulated_flow = cocoindex.Flow(
        file_source,
        extract_text_from_file, # Map operation
        text_splitter           # Split operation
    )
    
    simulated_flow.run()
    print("--- Flow Simulation Complete ---")
    print("\nNext steps would involve embedding these chunks and storing them in a vector database.")


# --- Main Execution Block ---
if __name__ == "__main__":
    print("Starting document ingestion pipeline demonstration...\n")

    # Create sample_documents directory if it doesn't exist (for local testing without the agent creating them)
    if not os.path.exists("./sample_documents"):
        os.makedirs("./sample_documents")
        print("Created './sample_documents/' directory for sample files.")
        # Create dummy files if they weren't created by the agent
        with open("./sample_documents/plain_text_example.txt", "w") as f: f.write("This is a test text file.")
        with open("./sample_documents/markdown_example.md", "w") as f: f.write("# Test MD")
        with open("./sample_documents/dummy_document.pdf", "w") as f: f.write("Dummy PDF")
        with open("./sample_documents/dummy_document.docx", "w") as f: f.write("Dummy DOCX")


    document_ingestion_flow()

    print("\nDemonstration finished.")
    print("To implement this for real, replace placeholder functions and CocoIndex modules")
    print("with actual implementations and install necessary libraries:")
    print("  pip install cocoindex pypdf python-docx")

```
